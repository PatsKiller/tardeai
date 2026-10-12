#!/usr/bin/env python3
"""Reconcile the existing workbook using explicit source inputs, without host probes.

Preview is the default. --apply also requires the inspected workbook SHA256.
Original inventory cells, decisions and sheets are retained; derived closure
fields and sheets never promote historical source/receipt labels to acceptance.
The workbook is the consumer of this operator-run report, not a runtime store.
openpyxl is needed only by the workbook adapter, not the pure reconciliation.
"""

from __future__ import annotations

import argparse
import csv
import fcntl
import hashlib
import json
import os
import subprocess
import threading
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

STATUSES = ("Completed", "In Progress", "Blocked", "Deferred", "Not Started")
MODIFICATIONS = {"Requires Refactoring Before Migration", "Requires Custom Development"}
CLASSIFICATIONS = MODIFICATIONS | {
    "Ready for Migration",
    "Candidate for Retirement",
    "Retain as Existing Cron Job",
    "Not Recommended for Migration",
}
UNKNOWN = "NOT_MEASURED"
DERIVED_SHEETS = {
    "Closure Dashboard",
    "Closure Items",
    "Modification Targets",
    "Closure Requirements",
    "Process Crosswalk",
    "Reconciliation Drift",
    "Reconciliation Evidence",
    "Contract Review",
    "Monitoring Coverage",
    "Owner Decisions",
    "Timer Source Gaps",
    "OpenClaw Crosswalk",
    "Runtime Observations",
    "Classification Review",
    "Coverage Resolution",
}
CUTOVER_FIELDS = ("reconciled_acceptance", "acceptance_basis", "acceptance_as_of")
_LOCKS = threading.local()


@contextmanager
def writer_lock(path):
    """Shared advisory lock with the two existing workbook generators.

    Cooperating writers fail before reading/writing. Raw filesystem writers that
    ignore this lock are outside the guarantee and must be serialized by the operator.
    """
    lock_path = Path(path).resolve().with_name(".CRON_INVENTORY_WRITER.lock")
    held = getattr(_LOCKS, "held", {})
    if str(lock_path) in held:
        yield
        return
    with lock_path.open("a+") as stream:
        fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        held[str(lock_path)] = True
        _LOCKS.held = held
        try:
            yield
        finally:
            del held[str(lock_path)]
            fcntl.flock(stream, fcntl.LOCK_UN)


REQUIREMENTS = {
    "classification": "Resolve inventory classification from current source and owner intent",
    "registry_declaration": "Confirm exact declared inventory registry ID exists in the source registry",
    "mapping": "Confirm business process, scheduler aliases, parent and exclusions",
    "disposition": "Record explicit owner/operator approved target disposition",
    "owner": "Obtain responsible owner's acceptance and dependency ownership",
    "source": "Review recursive entrypoint, environment and child closure at the integrated head",
    "refactor": "Accept required source changes with source diff and regression evidence",
    "testing": "Attach per-process functional, failure, timeout, idempotency and dependency tests",
    "deployment": "Verify authorized exact-CURRENT deployment and process code roots",
    "contract": "Review applicable health/output/receipt contract and its review provenance",
    "schedule": "Prove exact trigger, timezone, DST, market days and required natural fires",
    "outputs": "Correlate durable receipts with business outputs and downstream consumers",
    "duplicate_effects": "Inspect duplicate requests/effects and lock skips against the contract",
    "monitoring": "Prove failure to incident to notification and escalation for this process",
    "recovery": "Demonstrate terminal failure/retry policy and output recovery",
    "rollback": "Verify scoped rollback rehearsal and scheduler change packet if applicable",
    "documentation": "Review item-specific runbook, dependencies and recovery documentation",
    "cutover": "Review natural shadow, authorized canary and exact legacy-trigger disposition",
}
BASE_FIELDS = [
    "canonical_process_id",
    "crosswalk_status",
    "primary_work_status",
    "primary_status_basis",
    "approved_target_disposition",
    "refactor_status",
    "testing_acceptance",
    "deployment_acceptance",
    "contract_acceptance",
    "owner_acceptance",
    "cutover_acceptance",
    "production_readiness",
    "applicable_requirements_passed",
    "applicable_requirements_required",
    "requirement_coverage_pct",
    "closure_priority",
    "effort_class",
    "closure_risk",
    "closure_dependencies",
    "closure_blockers",
    "remaining_actions",
    "evidence_class",
    "evidence_paths",
    "observed_at",
    "tested_sha",
    "deployed_sha",
    "responsible_declared_owner",
    "ownership_confirmation",
    "next_action",
    "source_work_metadata",
    "lane_verdict",
    "decision_due_date",
    "natural_observation_date",
    "scoped_consolidation_deferral",
    "classification_proposal",
    "classification_review_status",
    "source_mapping_proposal",
    "source_declared_owner_candidate",
]


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def meaningful(value):
    text = str(value or "").strip()
    return bool(text and text.upper() not in {"NONE", "UNKNOWN", "N/A", "NOT_MEASURED"})


def contains_unknown(value):
    if isinstance(value, dict):
        return any(contains_unknown(v) for v in value.values())
    if isinstance(value, list):
        return any(contains_unknown(v) for v in value)
    return isinstance(value, str) and "UNKNOWN" in value.upper()


def full_deferral(row):
    decision = row.get("work_deferral_decision")
    if not isinstance(decision, dict):
        return False
    try:
        at = datetime.fromisoformat(str(decision.get("decided_at")))
    except ValueError:
        return False
    return (
        decision.get("scope") == "INVENTORY_CLOSURE"
        and decision.get("accepted") is True
        and at.tzinfo is not None
        and meaningful(decision.get("owner"))
        and meaningful(decision.get("evidence"))
    )


def reconcile(
    rows, *, registry, contracts, allowlist, completion, waves, codeonly, as_of, source_sha, deployed_sha, inputs
):
    """Conservative status partition; counts measure this review, never implied health.

    Acceptance is intentionally not imported from free-text notes. Collected item
    receipts require a separate explicit review before the acceptance fields move.
    """
    ids = [r.get("inv_id") for r in rows]
    if not all(ids) or len(set(ids)) != len(ids):
        raise ValueError("Inventory requires unique nonempty inv_id values")
    lanes = {r["lane_id"]: r for r in registry["lanes"]}
    if len(lanes) != len(registry["lanes"]):
        raise ValueError("Duplicate registry lane IDs")
    by_contract = {c["id"]: c for c in contracts["contracts"]}
    source = {r["lane_id"]: r for r in completion["completion119"]["rows"]}
    wave_by_id = defaultdict(list)
    for name, groups in waves.items():
        for group, values in groups.items():
            for inv_id in values["inv_ids"]:
                wave_by_id[inv_id].append(f"{name}:{group}")
    planning_ids = set(codeonly)
    blocked_reasons = allowlist.get("blocked_reasons", {})
    derived, matrix, crosswalk, drifts = [], [], [], []
    metadata_counts = Counter()
    for row in rows:
        inv_id = row["inv_id"]
        # Only an exact declared registry ID is canonical; name guesses remain proposals.
        declared_lane = str(row.get("registry_row") or "")
        lane = lanes.get(declared_lane)
        comp = source.get(declared_lane or str(row.get("job_name") or ""))
        if comp is None and not meaningful(declared_lane):
            comp = source.get(str(row.get("job_name") or ""))
        contract = by_contract.get(declared_lane)
        modification = row.get("classification") in MODIFICATIONS
        parent = str(row.get("x_parent") or "")
        owner = str(row.get("owner_stakeholder") or "")
        owner_unresolved = not meaningful(owner) or "UNKNOWN" in owner.upper()
        historical_blockers = str(row.get("blockers") or "")
        boundaries, findings, work = [], [], []
        if not lane:
            boundaries.append("Canonical registry/parent mapping requires reconciliation")
        if owner_unresolved:
            boundaries.append("Responsible owner unresolved; inherited parent owner is unconfirmed")
        if row.get("classification") not in CLASSIFICATIONS:
            boundaries.append("Classification and applicable closure scope require owner review")
        if declared_lane in blocked_reasons:
            boundaries.append("Current allowlist boundary: " + blocked_reasons[declared_lane])
        if comp and comp["source_result"] == "OWNER_OR_AUTHORITY_BLOCKED":
            boundaries.extend(comp["remaining"])
        if comp and comp.get("detail_result") == "DESIGN_ONLY":
            boundaries.extend(comp["remaining"])
        if lane and ((lane.get("scheduler") or {}).get("stage") in {"shadow", "canary", "cutover"}):
            if not contract or contract.get("status") != "REVIEWED" or contains_unknown(contract):
                boundaries.append("Current staged lane lacks a complete owner-reviewed contract")
        if lane and row.get("x_registry_state") and row["x_registry_state"] != lane.get("state"):
            findings.append("Workbook registry state differs from reviewed source")
        if historical_blockers and "unratified" in historical_blockers.lower():
            findings.append("Historical policy blocker must be reassessed against active source policy")
        if (
            row.get("migration_status") in {"Cut over", "Validated"}
            or row.get("validation_status") == "Natural-schedule proven"
        ):
            findings.append("Historical lifecycle/natural-proof label is not current acceptance")
        if meaningful(row.get("implemented")):
            work.append("Historical implementation annotation; acceptance unmeasured")
        if wave_by_id[inv_id]:
            work.extend(wave_by_id[inv_id])
        if comp:
            work.append("Completion119:" + comp["detail_result"])
        metadata_counts["wave_targets"] += int(modification and bool(wave_by_id[inv_id]))
        metadata_counts["completion_targets"] += int(modification and bool(comp))
        has_work_metadata = bool(wave_by_id[inv_id] or comp)
        metadata_counts["combined_targets"] += int(modification and has_work_metadata)
        metadata_counts["including_planning_targets"] += int(
            modification and (has_work_metadata or inv_id in planning_ids)
        )
        deferral = str(row.get("consolidation_approval") or "")
        deferred = full_deferral(row)
        scoped_deferral = deferral if deferral.startswith("Deferred") else "NONE"
        if scoped_deferral != "NONE" and not deferred:
            findings.append("Consolidation deferral does not defer full inventory closure; preserve its decision scope")
        if deferred:
            # The existing dated decision is explicit. It never proves other requirements.
            status, basis = "Deferred", json.dumps(row["work_deferral_decision"], sort_keys=True)
        elif boundaries:
            status, basis = "Blocked", "; ".join(boundaries)
        elif work or row.get("migration_status") in {"Shadow", "Retired", "Cut over", "Validated"}:
            status, basis = (
                "In Progress",
                "; ".join(work) or "Historical lifecycle evidence; closure remains unaccepted",
            )
        else:
            status, basis = "Not Started", "No implementation evidence under the closure contract"
        # Only declaration/mapping coverage can be measured from these inputs.
        passed = {
            "classification": row.get("classification") in CLASSIFICATIONS,
            "registry_declaration": lane is not None,
        }
        not_applicable = {}
        if row.get("classification") in CLASSIFICATIONS and not modification:
            not_applicable["refactor"] = (
                "Classification does not require source modification; recheck if disposition changes"
            )
        outside = row.get("classification") in {"Retain as Existing Cron Job", "Not Recommended for Migration"}
        if outside:
            for req in ("contract", "cutover"):
                not_applicable[req] = (
                    "Declared scope retains existing scheduler; production operations still require proof"
                )
        requirement_rows = []
        for name, action in REQUIREMENTS.items():
            state = "PROVEN" if passed.get(name) else UNKNOWN
            reason = ""
            if name in not_applicable:
                state, reason = "NOT_APPLICABLE_WITH_REASON", not_applicable[name]
            evidence = (
                inputs["workbook"]
                if name == "classification" and state == "PROVEN"
                else (inputs["registry"] if name == "registry_declaration" and state == "PROVEN" else "")
            )
            requirement_rows.append(
                {
                    "inv_id": inv_id,
                    "lane_id": lane["lane_id"] if lane else "",
                    "requirement": name,
                    "status": state,
                    "reason": reason,
                    "evidence": evidence,
                    "observed_at": as_of,
                    "tested_sha": UNKNOWN,
                    "deployed_sha": deployed_sha,
                    "collection_action": "" if state == "PROVEN" else reason or action,
                }
            )
        required = sum(r["status"] != "NOT_APPLICABLE_WITH_REASON" for r in requirement_rows)
        proven = sum(r["status"] == "PROVEN" for r in requirement_rows)
        matrix.extend(requirement_rows)
        actions = [r["collection_action"] for r in requirement_rows if r["status"] == UNKNOWN]
        if modification and meaningful(row.get("custom_work_required")):
            actions.insert(
                0,
                "Reassess historical custom work against current integrated source: "
                + str(row["custom_work_required"]),
            )
        if comp:
            actions = list(comp["remaining"]) + actions
        canonical = lane["lane_id"] if lane else "UNMAPPED:" + inv_id
        disposition = (
            deferral
            if deferred
            else (
                str(row.get("consolidation_approval"))
                if str(row.get("consolidation_approval") or "").startswith("Approved —")
                else "NOT_CONFIRMED; declared recommendation: " + str(row.get("classification") or "UNCLASSIFIED")
            )
        )
        d = {
            "inv_id": inv_id,
            "canonical_process_id": canonical,
            "crosswalk_status": "EXACT_REGISTRY; parent/business alias review pending" if lane else "UNRESOLVED",
            "primary_work_status": status,
            "primary_status_basis": basis,
            "approved_target_disposition": disposition,
            "refactor_status": "WORK_METADATA; ACCEPTANCE_NOT_MEASURED"
            if modification and has_work_metadata
            else (
                "PLANNING_ONLY; ACCEPTANCE_NOT_MEASURED"
                if modification and inv_id in planning_ids
                else (
                    UNKNOWN
                    if modification or row.get("classification") not in CLASSIFICATIONS
                    else "NOT_APPLICABLE_WITH_REASON"
                )
            ),
            "testing_acceptance": UNKNOWN,
            "deployment_acceptance": UNKNOWN,
            "contract_acceptance": "DECLARED_REVIEWED; PROVENANCE_AND_ACCEPTANCE_UNCONFIRMED"
            if contract and contract.get("status") == "REVIEWED"
            else ("DRAFT" if contract else "NOT_APPLICABLE_WITH_REASON" if outside else UNKNOWN),
            "owner_acceptance": UNKNOWN,
            "cutover_acceptance": "NOT_APPLICABLE_WITH_REASON" if outside else UNKNOWN,
            "production_readiness": UNKNOWN,
            "applicable_requirements_passed": proven,
            "applicable_requirements_required": required,
            "requirement_coverage_pct": round(100 * proven / required, 2),
            "closure_priority": "P1" if modification else row.get("migration_priority") or "UNASSIGNED",
            "effort_class": row.get("migration_complexity") or "NOT_ESTIMATED",
            "closure_risk": row.get("migration_risk") or "NOT_ASSESSED",
            "closure_dependencies": row.get("dependencies") or UNKNOWN,
            "closure_blockers": "; ".join(boundaries) or "No named boundary established; collection work remains",
            "remaining_actions": "; ".join(actions),
            "evidence_class": "SOURCE_AND_HISTORICAL_DECLARATIONS",
            "evidence_paths": "; ".join(inputs.values()),
            "observed_at": as_of,
            "tested_sha": UNKNOWN,
            "deployed_sha": deployed_sha,
            "responsible_declared_owner": owner or "UNKNOWN",
            "ownership_confirmation": UNKNOWN,
            "next_action": boundaries[0] if boundaries else actions[0],
            "source_work_metadata": "; ".join(work + (["code-only planning list"] if inv_id in planning_ids else []))
            or "NONE",
            "lane_verdict": "NO_GO" if not outside else "NOT_APPLICABLE_WITH_REASON: retained scheduler",
            "decision_due_date": "2026-10-12 (proposed owner review; not agreed)",
            "natural_observation_date": "NOT_SCHEDULED: requires reviewed contract and authorized exact deployment",
            "scoped_consolidation_deferral": scoped_deferral,
            "classification_proposal": "NONE",
            "classification_review_status": "NO_NEW_PROPOSAL",
            "source_mapping_proposal": "NONE",
            "source_declared_owner_candidate": "NONE",
        }
        derived.append(d)
        crosswalk.append(
            {
                "inv_id": inv_id,
                "canonical_process_id": canonical,
                "registry_lane_id": lane["lane_id"] if lane else "",
                "declared_registry_id": declared_lane,
                "parent": parent,
                "source": row.get("source"),
                "scheduler_id": inv_id,
                "workflow_id": row.get("proposed_n8n_workflow_id") or (contract or {}).get("workflow_id") or "",
                "source_scheduler": json.dumps((lane or {}).get("scheduler", {}), sort_keys=True),
                "historical_schedule": row.get("schedule_expr"),
                "contract_id": (contract or {}).get("id", ""),
                "declared_owner": owner,
                "dependencies": row.get("dependencies"),
                "test_evidence": UNKNOWN,
                "business_alias_acceptance": UNKNOWN,
            }
        )
        for finding in findings:
            drifts.append(
                {
                    "inv_id": inv_id,
                    "lane_id": declared_lane,
                    "finding": finding,
                    "evidence": inputs["workbook"],
                    "next_action": "Reconcile dated historical declaration with exact source/runtime evidence",
                }
            )
    targets = [d for r, d in zip(rows, derived) if r.get("classification") in MODIFICATIONS]
    partition = {s: sum(d["primary_work_status"] == s for d in derived) for s in STATUSES}
    assert sum(partition.values()) == len(rows)
    summary = {
        "as_of": as_of,
        "source_sha": source_sha,
        "deployed_sha": deployed_sha,
        "inventory_rows": len(rows),
        "classified_rows": sum(r.get("classification") in CLASSIFICATIONS for r in rows),
        "modification_targets": len(targets),
        "primary_status_partition": partition,
        "partition_basis": "Derived current boundaries and work evidence; preserves original lifecycle labels",
        "source_metadata_coverage": dict(metadata_counts),
        "targets_without_wave_or_completion_metadata": len(targets) - metadata_counts["combined_targets"],
        "targets_without_any_metadata_including_planning": len(targets) - metadata_counts["including_planning_targets"],
        "mapped_inventory_rows": sum(not d["canonical_process_id"].startswith("UNMAPPED:") for d in derived),
        "canonical_business_process_denominator": UNKNOWN,
        "exact_registry_ids_represented": len({d["registry_lane_id"] for d in crosswalk if d["registry_lane_id"]}),
        "classification_coverage_pct": round(
            100 * sum(r.get("classification") in CLASSIFICATIONS for r in rows) / len(rows), 2
        ),
        "inventory_closure_pct": UNKNOWN,
        "accepted_refactor_pct": UNKNOWN,
        "production_readiness_pct": UNKNOWN,
        "verified_completed_under_new_contract": partition["Completed"],
        "verified_completion_lower_bound_pct": round(100 * partition["Completed"] / len(rows), 2),
        "percentage_note": "Requirement coverage is checklist evidence coverage only. No combined maturity/readiness percentage.",
        "contracts_total": len(by_contract),
        "contracts_by_status": dict(Counter(c.get("status") for c in by_contract.values())),
        "all_targets_have_disposition_or_collection_action": all(d["next_action"] for d in targets),
        "historical_lifecycle": dict(Counter(r.get("migration_status") for r in rows)),
        "drift_findings": len(drifts),
        "drive_sync": "PENDING",
        "cutover": "BLOCKED",
    }
    return {
        "summary": summary,
        "items": derived,
        "targets": targets,
        "requirements": matrix,
        "crosswalk": crosswalk,
        "drift": drifts,
    }


def apply_coverage_proposals(result, packet, reference):
    """Attach exact source proposals without accepting mappings or owner decisions."""
    by_id = {r["inv_id"]: r for r in result["items"]}
    for proposal in packet["rows"]:
        item = by_id[proposal["inv_id"]]
        item["source_mapping_proposal"] = (
            proposal["source_mapping_status"] + ": " + json.dumps(proposal["source_backed_candidates"], sort_keys=True)
        )
        item["source_declared_owner_candidate"] = proposal["source_declared_owner"]
        item["next_action"] = proposal["remaining_action"]
        item["remaining_actions"] = proposal["remaining_action"] + "; " + item["remaining_actions"]
        item["evidence_paths"] += "; " + reference
    for proposal in packet["classification_proposals"]:
        item = by_id[proposal["inv_id"]]
        item["classification_proposal"] = proposal["proposed_classification"]
        item["classification_review_status"] = proposal["classification_proposal_status"]
        item["source_declared_owner_candidate"] = proposal["source_declared_owner"]
        item["remaining_actions"] = (
            "Review source-backed classification: " + proposal["rationale"] + "; " + item["remaining_actions"]
        )
    result["summary"]["coverage_reconciliation"] = packet["counts"]


def write_artifacts(result, output):
    output.mkdir(parents=True, exist_ok=True)
    (output / "reconciliation.json").write_text(json.dumps(result, indent=2) + "\n")
    (output / "summary.json").write_text(json.dumps(result["summary"], indent=2) + "\n")
    for key in ("items", "targets", "requirements", "crosswalk", "drift"):
        values = result[key]
        if not values:
            continue
        with (output / (key + ".tsv")).open("w", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(values[0]), delimiter="\t")
            writer.writeheader()
            writer.writerows(values)


def sheet_table(wb, name, headers, rows):
    from openpyxl.styles import Alignment, Font, PatternFill

    if name in wb:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203864")
    for cells in ws:
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row > 1:
                cell.font = Font(name="Arial", size=10)
        if cells[0].row > 1:
            ws.row_dimensions[cells[0].row].height = 45
    for cell in ws[1]:
        ws.column_dimensions[cell.column_letter].width = min(65, max(20, len(str(cell.value)) + 4))
    return ws


def original_cells(wb):
    """Preservation receipt covers original values, links, comments and styles."""
    return {
        s.title: [
            [
                (
                    c.coordinate,
                    c.value,
                    c.hyperlink.target if c.hyperlink else None,
                    c.comment.text if c.comment else None,
                    list(c._style) if c.has_style else None,
                )
                for c in row
                if not (
                    (s.title == "Inventory" and s.cell(1, c.column).value in BASE_FIELDS)
                    or (s.title == "Cutover_Tracker" and s.cell(1, c.column).value in CUTOVER_FIELDS)
                )
            ]
            for row in s
        ]
        for s in wb
        if s.title not in DERIVED_SHEETS
    }


def verify_preserved(wb, original):
    for name, rows in original.items():
        for row in rows:
            for coord, value, link, comment, style in row:
                c = wb[name][coord]
                current = (
                    c.value,
                    c.hyperlink.target if c.hyperlink else None,
                    c.comment.text if c.comment else None,
                    list(c._style) if c.has_style else None,
                )
                if current != (value, link, comment, style):
                    raise ValueError(f"Original cell changed: {name}!{coord}")


def workbook_candidate(wb, result, output):
    inventory = wb["Inventory"]
    headers = {c.value: c.column for c in inventory[1]}
    for field in BASE_FIELDS:
        if field not in headers:
            headers[field] = inventory.max_column + 1
            inventory.cell(1, headers[field], field)
    for index, item in enumerate(result["items"], 2):
        for field in BASE_FIELDS:
            inventory.cell(index, headers[field], item[field])
    inventory.auto_filter.ref = inventory.dimensions
    for key, name in (
        ("items", "Closure Items"),
        ("targets", "Modification Targets"),
        ("requirements", "Closure Requirements"),
        ("crosswalk", "Process Crosswalk"),
        ("drift", "Reconciliation Drift"),
    ):
        values = result[key]
        if values:
            h = list(values[0])
            sheet_table(wb, name, h, [[v[k] for k in h] for v in values])
    summary = result["summary"]
    sheet_table(
        wb,
        "Closure Dashboard",
        ["metric", "value", "basis"],
        [
            [k, json.dumps(v) if isinstance(v, (dict, list)) else v, "Reconciliation at " + summary["as_of"]]
            for k, v in summary.items()
        ],
    )
    dashboard = wb["Closure Dashboard"]
    dashboard.column_dimensions["A"].width = 45
    dashboard.column_dimensions["B"].width = 95
    dashboard.column_dimensions["C"].width = 45
    for status, count in summary["primary_status_partition"].items():
        dashboard.append(["Primary work status: " + status, count, "Current closure contract; historical labels retained"])
    sheet_table(
        wb,
        "Reconciliation Evidence",
        ["path", "sha256"],
        [[str(p), sha256(p.read_bytes())] for p in sorted(output.glob("*.json"))],
    )
    # Original tracker rows remain a historical audit trail. New acceptance is explicit.
    tracker = wb["Tracker"]
    tracker.append(
        ["Reconciliation: historical rows above do not measure closure; use Closure Dashboard", summary["as_of"]]
    )
    for status, count in summary["primary_status_partition"].items():
        tracker.append(["Primary closure status: " + status, count])
    for field in ("inventory_closure_pct", "accepted_refactor_pct", "production_readiness_pct"):
        tracker.append([field, summary[field]])
    readme = wb["Read Me"]
    readme.append(
        [
            "Closure Dashboard supersedes historical lifecycle totals for acceptance. Closure Items retains every inventory row; Process Crosswalk uses exact IDs and explicitly unresolved parents/aliases. Requirement coverage is not production readiness. Original cells and operator decisions are preserved. Re-run scripts/reconcile_migration_inventory.py with explicit source inputs; rebuilding only build_workbook.py omits this review. Source artifacts: "
            + str(output)
        ]
    )
    if "Cutover_Tracker" in wb:
        ws = wb["Cutover_Tracker"]
        columns = {c.value: c.column for c in ws[1]}
        for name in CUTOVER_FIELDS:
            if name not in columns:
                columns[name] = ws.max_column + 1
                ws.cell(1, columns[name], name)
        for r in range(2, ws.max_row + 1):
            ws.cell(r, columns[CUTOVER_FIELDS[0]], "NO_GO")
            ws.cell(
                r,
                columns[CUTOVER_FIELDS[1]],
                "Current exact-head source/deployment/natural/output/monitoring acceptance is NOT_MEASURED; legacy stage retained",
            )
            ws.cell(r, columns[CUTOVER_FIELDS[2]], summary["as_of"])


def publish_candidate(path, candidate, *, expected_sha, original_bytes):
    """Hash-checked publication serialized with cooperating workbook writers."""
    with writer_lock(path):
        return _publish_locked(path, candidate, expected_sha=expected_sha, original_bytes=original_bytes)


def _publish_locked(path, candidate, *, expected_sha, original_bytes):
    if sha256(path.read_bytes()) != expected_sha or sha256(original_bytes) != expected_sha:
        raise ValueError("Workbook changed after inspection; candidate retained, publication refused")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    backup = path.with_name(path.name + ".pre-reconciliation-" + stamp)
    with backup.open("xb") as stream:
        stream.write(original_bytes)
        stream.flush()
        os.fsync(stream.fileno())
    temp = path.with_name("." + path.stem + ".reconciliation-" + stamp + ".xlsx")
    with temp.open("xb") as stream:
        stream.write(candidate.read_bytes())
        stream.flush()
        os.fsync(stream.fileno())
    if sha256(path.read_bytes()) != expected_sha:
        raise ValueError("Workbook changed during publication; backup and candidate retained")
    os.replace(temp, path)
    return {
        "backup": str(backup),
        "before_sha256": expected_sha,
        "after_sha256": sha256(path.read_bytes()),
        "original_cells_preserved": True,
        "writer_lock": "shared advisory .CRON_INVENTORY_WRITER.lock; all registered writers serialized",
    }


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--workbook", required=True, type=Path)
    ap.add_argument("--inventory-dir", required=True, type=Path)
    ap.add_argument("--source-root", required=True, type=Path)
    ap.add_argument("--source-sha", required=True)
    ap.add_argument("--deployed-sha", required=True)
    ap.add_argument("--output", required=True, type=Path)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--expected-sha")
    ap.add_argument("--supplement", type=Path, help="sanitized runtime-contracts packet directory")
    args = ap.parse_args(argv)
    if args.apply and not args.expected_sha:
        ap.error("--apply requires --expected-sha from workbook inspection")
    with writer_lock(args.workbook):
        return build(args)


def build(args):
    from openpyxl import load_workbook

    actual_sha = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=args.source_root, text=True).strip()
    if actual_sha != args.source_sha:
        raise ValueError("Declared source SHA differs from source-root HEAD")

    original_bytes = args.workbook.read_bytes()
    if args.expected_sha and sha256(original_bytes) != args.expected_sha:
        raise ValueError("Workbook changed after inspection")
    wb = load_workbook(args.workbook)
    protected = original_cells(wb)
    headers = [c.value for c in wb["Inventory"][1]]
    rows = [dict(zip(headers, values)) for values in list(wb["Inventory"].values)[1:]]
    files = {
        "workbook": args.workbook,
        "registry": args.source_root / "config/lane_registry.json",
        "contracts": args.source_root / "config/n8n_health_contracts.json",
        "allowlist": args.source_root / "config/n8n_run_allowlist.json",
        "completion": args.inventory_dir / "source-progress/CODEX_SOURCE_PROGRESS_20261011.json",
    }
    load = lambda p: json.loads(p.read_text())
    provenance = {k: str(p) + " #sha256=" + sha256(p.read_bytes()) for k, p in files.items()}
    waves = {f"wave{i}": load(args.inventory_dir / f"refactor_wave{i}.json") for i in (1, 2, 3)}
    provenance.update(
        {
            f"wave{i}": str(args.inventory_dir / f"refactor_wave{i}.json")
            + " #sha256="
            + sha256((args.inventory_dir / f"refactor_wave{i}.json").read_bytes())
            for i in (1, 2, 3)
        }
    )
    provenance["codeonly"] = (
        str(args.inventory_dir / "refactor_codeonly.json")
        + " #sha256="
        + sha256((args.inventory_dir / "refactor_codeonly.json").read_bytes())
    )
    result = reconcile(
        rows,
        registry=load(files["registry"]),
        contracts=load(files["contracts"]),
        allowlist=load(files["allowlist"]),
        completion=load(files["completion"]),
        waves=waves,
        codeonly=load(args.inventory_dir / "refactor_codeonly.json"),
        as_of=datetime.now(timezone.utc).isoformat(),
        source_sha=args.source_sha,
        deployed_sha=args.deployed_sha,
        inputs=provenance,
    )
    if args.supplement:
        runtime = load(args.supplement / "runtime_observations.json")
        contracts = load(args.supplement / "contract_reconciliation.json")
        result["summary"]["runtime_observations"] = {"as_of": runtime["as_of"], **runtime["metrics"]}
        result["summary"]["contract_reconciliation_counts"] = contracts["counts"]
        provenance["runtime_packet"] = (
            str(args.supplement / "runtime_observations.json")
            + " #sha256="
            + sha256((args.supplement / "runtime_observations.json").read_bytes())
        )
        coverage_path = args.supplement / "inventory_coverage_resolution.json"
        coverage = load(coverage_path) if coverage_path.exists() else None
        if coverage:
            provenance["coverage_packet"] = str(coverage_path) + " #sha256=" + sha256(coverage_path.read_bytes())
            apply_coverage_proposals(result, coverage, provenance["coverage_packet"])
    write_artifacts(result, args.output)
    workbook_candidate(wb, result, args.output)
    if args.supplement:
        if coverage:
            for key, title in (("rows", "Coverage Resolution"), ("classification_proposals", "Classification Review")):
                headers = sorted({k for r in coverage[key] for k in r})
                sheet_table(
                    wb,
                    title,
                    headers,
                    [
                        [json.dumps(r.get(k)) if isinstance(r.get(k), (list, dict)) else r.get(k) for k in headers]
                        for r in coverage[key]
                    ],
                )
        for filename, title in (
            ("monitoring_coverage.tsv", "Monitoring Coverage"),
            ("owner_decisions.tsv", "Owner Decisions"),
            ("timer_source_gaps.tsv", "Timer Source Gaps"),
            ("openclaw_crosswalk.tsv", "OpenClaw Crosswalk"),
        ):
            with (args.supplement / filename).open() as stream:
                values = list(csv.reader(stream, delimiter="\t"))
            sheet_table(wb, title, values[0], values[1:])
        h = list(contracts["rows"][0])
        sheet_table(
            wb,
            "Contract Review",
            h,
            [
                [json.dumps(r.get(k)) if isinstance(r.get(k), (list, dict)) else r.get(k) for k in h]
                for r in contracts["rows"]
            ],
        )
        sheet_table(
            wb,
            "Runtime Observations",
            ["section", "evidence"],
            [[k, json.dumps(v)] for k, v in runtime.items() if k != "ledger"],
        )
    # Place the authoritative progress view first without altering original cells.
    dashboard = wb["Closure Dashboard"]
    wb._sheets.remove(dashboard)
    wb._sheets.insert(0, dashboard)
    candidate = args.output / "CRON_INVENTORY_RECONCILED.xlsx"
    wb.save(candidate)
    wb.close()
    check = load_workbook(candidate)
    verify_preserved(check, protected)
    assert check["Closure Items"].max_row == len(rows) + 1
    assert check["Modification Targets"].max_row == len(result["targets"]) + 1
    check.close()
    receipt = {
        "as_of": result["summary"]["as_of"],
        "workbook": str(args.workbook),
        "candidate": str(candidate),
        "source_inputs": provenance,
        "candidate_sha256": sha256(candidate.read_bytes()),
        "original_cells_preserved": True,
        "published": args.apply,
        "drive_sync": "PENDING",
    }
    if args.apply:
        receipt.update(
            publish_candidate(args.workbook, candidate, expected_sha=args.expected_sha, original_bytes=original_bytes)
        )
    (args.output / "workbook-reconciliation-receipt.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(
        json.dumps(
            {**result["summary"], "receipt": str(args.output / "workbook-reconciliation-receipt.json")}, indent=2
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
