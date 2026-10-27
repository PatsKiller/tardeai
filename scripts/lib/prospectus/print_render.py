"""Print-fidelity renders: the PDF and the Word document are made from the SAME HTML report the reader sees.

Operator 2026-10-27: "the HTML prospectus looks great. The PDF and the Word document doesn't … so the Word document
and the PDF look identical."

* PDF  = headless Chromium (Playwright, already pinned in requirements.txt) printing the light-theme HTML report with
         a print stylesheet: Letter/A4, no split charts / KPI rows / table rows, repeated table headers, and a running
         header + footer (symbol, version date, page x / y).
* DOCX = an editable Word document converted by LibreOffice (``soffice --headless``, isolated temp profile) from a
         Word-oriented variant of the same HTML: literal light-theme colours, KPI chips as a shaded table, sections as
         shaded cards, and the charts as PNGs that Chromium rasterised from the very same SVGs. python-docx then sets
         margins, a header/footer with page numbers, repeated table header rows and the document title.

Both read the one presentation model through ``render.html`` and strip every citation marker (footnote numbers,
``[F7]`` ids, ``data-fact-id`` attributes, per-section "Sources" and the Sources appendix): the downloads never print
fact ids or footnotes. Anything that fails (no Chromium, no LibreOffice, timeout, a crash) falls back to the previous
renderer for that one format (``render.pdf`` WeasyPrint / ``render.docx`` python-docx) and logs the reason.

Runs only in the build job (``pipeline.persist_doc`` → ``render.render_all``), never in an API request. One render
at a time (a host lock), a hard timeout per tool, JavaScript off and every network request aborted in the print page.
READ_ONLY_ADVISORY: rendering only (MBI_BEHAVIOR=0).
"""
from __future__ import annotations

import base64
from contextlib import contextmanager
import fcntl
import io
import json
import logging
import os
from pathlib import Path
import re
import shutil
import signal
import struct
import subprocess
import sys
import tempfile
import time
from html import escape
from typing import Any, Callable, Iterator

from . import charts as ch

LOG = logging.getLogger("prospectus.print")
WORKER = Path(__file__).with_name("print_worker.py")

DEFAULTS: dict[str, Any] = {
    "enabled": True,
    "paper": "Letter",              # "Letter" | "A4"
    "margin_in": {"top": 0.75, "bottom": 0.7, "left": 0.6, "right": 0.6},
    "chromium_timeout_s": 90,
    "soffice_timeout_s": 120,
    "lock_wait_s": 300,
    "chart_scale": 2,               # device pixels per CSS pixel for chart PNGs in Word
    "chart_width_px": 640,          # matches charts.W (the SVG viewBox width)
    "docx_font": "Calibri",
    "docx_image_width_px": 672,     # 7.0in at 96 px/in: the Letter text column less a little
    "soffice": "soffice",
    "chromium_args": [],            # Playwright launches its bundled headless shell without the setuid sandbox; JS
                                    # is off and all network aborted, so only our own static HTML is ever parsed
}

LIGHT = ch.TOKENS["light"]

# ── citations: the downloads never print fact ids or footnotes ──────────────────────────────────────────────────────
_CITE_PATTERNS = (
    re.compile(r'\s*<sup class="pc-fn">.*?</sup>', re.S),
    re.compile(r'<details class="pc-srcs">.*?</details>', re.S),
    re.compile(r'<section id="sources"[^>]*>.*?</section>', re.S),
    re.compile(r"<script\b.*?</script>", re.S),
    re.compile(r'\sdata-fact-id="[^"]*"'),
    re.compile(r"\s*\[F\d+(?:\s*,\s*F\d+)*\]"),
)


def strip_citations(html: str) -> str:
    for rx in _CITE_PATTERNS:
        html = rx.sub("", html)
    return html


CONFIG_REL = "config/prospectus_print.json"


def settings(cfg: dict | None = None) -> dict:
    """DEFAULTS overlaid with ``config/prospectus_print.json`` (or ``cfg["print"]`` when a config is passed)."""
    if cfg is None:
        try:
            from . import ROOT  # noqa: PLC0415
            cfg = {"print": json.loads((ROOT / CONFIG_REL).read_text(encoding="utf-8"))}
        except (OSError, ValueError):
            cfg = {}
    out = json.loads(json.dumps(DEFAULTS))
    for k, v in ((cfg or {}).get("print") or {}).items():
        if str(k).startswith("_") or k == "schema":
            continue
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k].update(v)
        else:
            out[k] = v
    return out


def version_date(doc: dict) -> str:
    v = str(doc.get("version") or doc.get("generated_at") or "")
    m = re.match(r"^(\d{4})-?(\d{2})-?(\d{2})", v)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else v[:10]


def _title(doc: dict) -> tuple[str, str]:
    from . import presentation as pr  # noqa: PLC0415
    sym = str(doc.get("symbol") or "")
    try:
        name = str(pr.present(doc).get("display_name") or doc.get("name") or "")
    except Exception:  # noqa: BLE001
        name = str(doc.get("name") or "")
    return sym, (name if name.upper() != sym.upper() else "")


# ── PDF: the HTML report with a print stylesheet ────────────────────────────────────────────────────────────────────
PRINT_CSS = """
@page{size:%(paper)s}
html,body{background:#fff!important}
*{-webkit-print-color-adjust:exact;print-color-adjust:exact}
body{font-size:10.5pt;line-height:1.45}
main{max-width:none;margin:0;padding:0}
header.pc-head{break-inside:avoid;break-after:avoid}
section,figure{-webkit-box-decoration-break:clone;box-decoration-break:clone}
section{margin:10px 0;padding:10px 12px}
h1,h2,h3,figcaption,section>h2{break-after:avoid;break-inside:avoid}
figure,svg,img,.pc-kpis,.pc-kpi,tr,li,.pc-cov,.pc-partial,.pc-note{break-inside:avoid}
p{orphans:3;widows:3}
.pc-tw{overflow:visible}
table.pc-table{break-inside:auto}
table.pc-table thead{display:table-header-group}
table.pc-table tr{break-inside:avoid;break-after:auto}
.pc-charts{display:block}
.pc-charts figure,figure{margin:10px 0}
a{color:inherit;text-decoration:none}
"""


def print_html(doc: dict, *, paper: str = "Letter") -> str:
    """The reader's HTML report (same model, same order) in the light print theme, citations removed."""
    from . import render  # noqa: PLC0415 - render imports this module lazily
    html = strip_citations(render.html(doc, theme="light", print_mode=True))
    return html.replace("</style>", PRINT_CSS % {"paper": paper} + "</style>", 1)


def _hf_templates(doc: dict, s: dict) -> tuple[str, str]:
    sym, name = _title(doc)
    m = s["margin_in"]
    pad = f"padding:0 {m['right']}in 0 {m['left']}in"
    style = (f"font:8px system-ui,-apple-system,'Segoe UI',Roboto,sans-serif;color:{LIGHT['muted']};width:100%;{pad};"
             "display:flex;justify-content:space-between;align-items:center;-webkit-print-color-adjust:exact")
    head = (f'<div style="{style}"><span><b style="color:{LIGHT["text"]}">{escape(sym)}</b>'
            f'{(" · " + escape(name)) if name else ""}</span><span>Prospectus · version {escape(version_date(doc))}</span></div>')
    foot = (f'<div style="{style}"><span>{"Holding" if doc.get("scope") == "holding" else "Watch list"} · '
            f'generated {escape(str(doc.get("generated_at") or "")[:16])} UTC · advisory only</span>'
            '<span>Page <span class="pageNumber"></span> / <span class="totalPages"></span></span></div>')
    return head, foot


# ── DOCX: a Word-oriented HTML variant of the same report ───────────────────────────────────────────────────────────
def _resolve_vars(text: str) -> str:
    tok = dict(LIGHT)
    for i, c in enumerate(LIGHT["chart"]):
        tok[f"chart-{i + 1}"] = c
    return re.sub(r"var\(--([\w-]+)\)", lambda m: tok.get(m.group(1), "#000000"), text)


WORD_CSS = """
body{font-family:Calibri,'Segoe UI',Arial,sans-serif;font-size:10.5pt;color:%(text)s}
h1{font-size:20pt;color:%(text)s;margin:0 0 2pt 0}
h1 small,h2 small{font-size:11pt;color:%(muted)s;font-weight:normal}
h2{font-size:13pt;color:%(text)s;margin:0 0 4pt 0}
p{margin:0 0 5pt 0}
li{margin:0 0 2pt 0}
.pc-verdict{font-size:12pt;color:%(text2)s}
.pc-meta,.pc-note,.pc-foot,.pc-cap{font-size:8.5pt;color:%(muted)s}
.pc-mute{color:%(muted)s;font-style:italic}
.pc-cap{font-size:9.5pt;color:%(text2)s;margin:6pt 0 2pt 0}
th{font-size:8pt;color:%(muted)s;text-align:left}
td{font-size:9.5pt;vertical-align:top}
"""


def _png_size(data: bytes) -> tuple[int, int]:
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return struct.unpack(">II", data[16:24])
    return (0, 0)


def word_html(doc: dict, chart_pngs: dict[str, bytes] | None = None, *, width_px: int = 624) -> str:
    """Word-oriented HTML from the same report: LibreOffice reads tables, inline styles, bgcolor and images, not
    flex/grid/CSS variables. ``chart_pngs`` maps chart id → PNG bytes (embedded as data: URIs so they are stored in
    the .docx, not linked)."""
    html = print_html(doc)
    chart_pngs = chart_pngs or {}
    L = LIGHT
    card_open = (f'<table width="100%" cellpadding="8" cellspacing="0" bgcolor="{L["surface"]}" '
                 f'style="border:1px solid {L["border"]};background:{L["surface"]};margin:0 0 8pt 0"><tr><td>')
    card_close = "</td></tr></table><p style=\"margin:0;font-size:4pt\">&nbsp;</p>"

    def figure(m: re.Match) -> str:
        cid, cap = m.group(1), m.group(2)
        data = chart_pngs.get(cid)
        if not data:
            return ""
        w, h = _png_size(data)
        hh = int(round(width_px * h / w)) if w else int(width_px * 0.4)
        return (f'<p class="pc-cap"><b>{cap}</b></p>'
                f'<p><img src="{data_uri(data)}" width="{width_px}" height="{hh}" alt="{cap}"></p>')

    html = re.sub(r'<figure id="chart-([\w-]+)"><figcaption>(.*?)</figcaption>.*?</figure>', figure, html, flags=re.S)
    html = re.sub(r'<div class="pc-charts">(.*?)</div>', r"\1", html, flags=re.S)

    def kpis(m: re.Match) -> str:
        cells = re.findall(r'<div class="pc-kpi"><span>(.*?)</span><b>(.*?)</b></div>', m.group(1), re.S)
        if not cells:
            return ""
        rows, per = [], 4
        for i in range(0, len(cells), per):
            chunk = cells[i:i + per]
            tds = "".join(f'<td bgcolor="{L["surface2"]}" style="background:{L["surface2"]};border:2px solid {L["surface"]}">'
                          f'<font size="1" color="{L["muted"]}">{k}</font><br><b>{v}</b></td>' for k, v in chunk)
            rows.append(f"<tr>{tds}</tr>")
        return f'<table cellpadding="5" cellspacing="0" style="margin:2pt 0 6pt 0">{"".join(rows)}</table>'

    html = re.sub(r'<div class="pc-kpis">(.*?</div>)</div>', kpis, html, flags=re.S)
    html = re.sub(r'<div class="pc-tw">(.*?)</div>', r"\1", html, flags=re.S)

    def table(m: re.Match) -> str:
        head, body = m.group(1), m.group(2)
        head = re.sub(r"<th>(.*?)</th>", lambda x: (f'<th bgcolor="{L["surface2"]}" style="background:{L["surface2"]};'
                                                     f'border-bottom:1px solid {L["border"]}"><font size="1" color="{L["muted"]}">'
                                                     f'{x.group(1).upper()}</font></th>'), head, flags=re.S)
        body = re.sub(r'<tr class="pc-hl">(.*?)</tr>', lambda x: "<tr>" + re.sub(
            r"<td>(.*?)</td>", lambda y: f'<td bgcolor="{L["surface2"]}" style="background:{L["surface2"]}"><b>{y.group(1)}</b></td>',
            x.group(1), flags=re.S) + "</tr>", body, flags=re.S)
        body = re.sub(r"<td>", f'<td bgcolor="#ffffff" style="background:#ffffff;border-top:1px solid {L["border"]}">', body)
        return (f'<table width="100%" cellpadding="4" cellspacing="0" bgcolor="#ffffff" style="border:1px solid {L["border"]};'
                f'background:#ffffff"><thead>{head}</thead><tbody>{body}</tbody></table>')

    html = re.sub(r'<table class="pc-table"><thead>(.*?)</thead><tbody>(.*?)</tbody></table>', table, html, flags=re.S)
    tone = {"used": L["success"], "stale": L["warning"], "empty": L["warning"], "failed": L["danger"],
            "not_applicable": L["muted"]}
    html = re.sub(r'<span class="pc-st pc-st-(\w+)"[^>]*>(.*?)</span>',
                  lambda m: f'<font color="{tone.get(m.group(1), L["text2"])}">{m.group(2)}</font>', html, flags=re.S)
    def meta(m: re.Match) -> str:
        inner = re.sub(r'<span class="pc-badge[^"]*">(.*?)</span>', lambda b: f'<b><font color="{L["info"]}">{b.group(1)}</font></b>',
                       m.group(1), flags=re.S)
        items = re.findall(r"<b><font.*?</font></b>|<span>.*?</span>", inner, re.S)
        return '<p class="pc-meta">' + " · ".join(re.sub(r"</?span>", "", i) for i in items) + "</p>"

    html = re.sub(r'<div class="pc-meta">(.*?)</div>', meta, html, flags=re.S)
    html = re.sub(r'<header class="pc-head">(.*?)</header>',
                  lambda m: m.group(1) + f'<hr style="border:0;border-top:2px solid {L["info"]};height:0;margin:2pt 0 8pt 0">',
                  html, flags=re.S)
    html = re.sub(r'<p class="pc-partial">(.*?)</p>',
                  lambda m: (f'<table width="100%" cellpadding="6" cellspacing="0"><tr><td bgcolor="{L["warning"]}" width="4"></td>'
                             f'<td bgcolor="{L["surface2"]}" style="background:{L["surface2"]}">{m.group(1)}</td></tr></table>'),
                  html, flags=re.S)
    html = re.sub(r'<section id="[^"]*"(?: class="[^"]*")?>', card_open, html)
    html = html.replace("</section>", card_close)
    html = re.sub(r"<style>.*?</style>", "<style>" + WORD_CSS % L + "</style>", html, count=1, flags=re.S)
    html = re.sub(r'<html lang="en"[^>]*>', '<html lang="en">', html)
    html = re.sub(r"<main>|</main>", "", html)
    return _resolve_vars(html)


# ── process plumbing ───────────────────────────────────────────────────────────────────────────────────────────────
@contextmanager
def _host_lock(wait_s: float) -> Iterator[bool]:
    """Concurrency 1 across processes: one print render at a time on this host."""
    path = Path(tempfile.gettempdir()) / "prospectus_print.lock"
    fh = open(path, "a+")
    got = False
    deadline = time.time() + max(0.0, float(wait_s))
    try:
        while True:
            try:
                fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
                got = True
                break
            except OSError:
                if time.time() >= deadline:
                    break
                time.sleep(0.5)
        yield got
    finally:
        if got:
            fcntl.flock(fh, fcntl.LOCK_UN)
        fh.close()


def _run(cmd: list[str], timeout_s: float, *, env: dict | None = None, cwd: str | None = None) -> subprocess.CompletedProcess:
    """Run in its own process group; on timeout the whole group (soffice.bin, chromium children) is killed."""
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, env=env, cwd=cwd,
                            start_new_session=True)
    try:
        out, err = proc.communicate(timeout=timeout_s)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(proc.pid, signal.SIGKILL)
        except OSError:
            pass
        proc.communicate()
        raise
    return subprocess.CompletedProcess(cmd, proc.returncode, out, err)


def chromium_available() -> str | None:
    """None when usable, else the reason."""
    try:
        import playwright  # noqa: F401, PLC0415
    except ImportError:
        return "python playwright not installed"
    return None


def soffice_path(s: dict) -> str | None:
    return shutil.which(str(s.get("soffice") or "soffice"))


def browser_env() -> dict:
    """Process env for the worker: Playwright finds its browsers under $HOME/.cache/ms-playwright, so a job whose
    HOME is redirected (a sandboxed lane, pytest) points PLAYWRIGHT_BROWSERS_PATH at the account's real cache."""
    env = dict(os.environ)
    if not env.get("PLAYWRIGHT_BROWSERS_PATH") and not (Path(env.get("HOME") or "~").expanduser() / ".cache/ms-playwright").is_dir():
        try:
            import pwd  # noqa: PLC0415
            real = Path(pwd.getpwuid(os.getuid()).pw_dir) / ".cache/ms-playwright"
            if real.is_dir():
                env["PLAYWRIGHT_BROWSERS_PATH"] = str(real)
        except (ImportError, KeyError):
            pass
    return env


def run_chromium(job: dict, timeout_s: float) -> dict:
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, dir=job["_dir"]) as fh:
        json.dump({k: v for k, v in job.items() if not k.startswith("_")}, fh)
    try:
        cp = _run([sys.executable, str(WORKER), fh.name], timeout_s, env=browser_env())
    except subprocess.TimeoutExpired:
        return {"ok": False, "error": f"chromium timeout after {timeout_s}s"}
    line = (cp.stdout or "").strip().splitlines()[-1:] or [""]
    try:
        res = json.loads(line[0])
    except ValueError:
        res = {"ok": False, "error": f"worker exit {cp.returncode}: {(cp.stderr or '')[-300:]}"}
    return res


def soffice_convert(src: Path, out_dir: Path, target: str, s: dict, *, infilter: str | None = None) -> Path:
    """LibreOffice headless conversion in an isolated, throwaway profile. Raises on failure."""
    exe = soffice_path(s)
    if not exe:
        raise RuntimeError("soffice not found")
    profile = out_dir / "lo-profile"
    profile.mkdir(exist_ok=True)
    cmd = [exe, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless", "--norestore", "--nolockcheck",
           "--nodefault", "--nologo"]
    if infilter:
        cmd.append(f"--infilter={infilter}")
    cmd += ["--convert-to", target, "--outdir", str(out_dir), str(src)]
    env = dict(os.environ, HOME=str(out_dir), SAL_USE_VCLPLUGIN="svp")
    cp = _run(cmd, float(s["soffice_timeout_s"]), env=env, cwd=str(out_dir))
    ext = target.split(":")[0]
    out = out_dir / f"{src.stem}.{ext}"
    if not out.exists() or out.stat().st_size == 0:
        raise RuntimeError(f"soffice exit {cp.returncode}: {(cp.stderr or cp.stdout or '')[-300:]}")
    return out


# ── DOCX post-processing (python-docx) ──────────────────────────────────────────────────────────────────────────────
def _field(para: Any, instr: str, style: Callable[[Any], None]) -> None:
    """A PAGE / NUMPAGES field, each part in its own run carrying the footer's run style."""
    from docx.oxml import OxmlElement  # noqa: PLC0415
    from docx.oxml.ns import qn  # noqa: PLC0415
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        run = para.add_run()
        style(run)
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def _set_fonts(root: Any, font: str) -> None:
    """LibreOffice keeps the CSS font list verbatim ("Calibri;Segoe UI;…") or its own defaults (Liberation Serif):
    every rFonts becomes the one Word font, so Word, LibreOffice and Pages all show the same face."""
    from docx.oxml.ns import qn  # noqa: PLC0415
    for el in root.iter(qn("w:rFonts")):
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            el.set(qn(attr), font)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if el.get(qn(attr)) is not None:
                del el.attrib[qn(attr)]


def postprocess_docx(data: bytes, doc: dict, s: dict) -> bytes:
    from docx import Document  # noqa: PLC0415
    from docx.oxml import OxmlElement  # noqa: PLC0415
    from docx.oxml.ns import qn  # noqa: PLC0415
    from docx.shared import Inches, Pt, RGBColor  # noqa: PLC0415

    d = Document(io.BytesIO(data))
    sym, name = _title(doc)
    m = s["margin_in"]
    a4 = str(s.get("paper") or "Letter").upper() == "A4"
    muted = RGBColor.from_string(LIGHT["muted"].lstrip("#").upper())
    for sec in d.sections:
        sec.page_width, sec.page_height = (Inches(8.27), Inches(11.69)) if a4 else (Inches(8.5), Inches(11))
        sec.top_margin, sec.bottom_margin = Inches(m["top"]), Inches(m["bottom"])
        sec.left_margin, sec.right_margin = Inches(m["left"]), Inches(m["right"])
        sec.header_distance = sec.footer_distance = Inches(0.35)
        text_w = sec.page_width - sec.left_margin - sec.right_margin
        tab = text_w
        hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        hp.text = ""
        r = hp.add_run(sym)
        r.bold = True
        r.font.size = Pt(8)
        r = hp.add_run((f" · {name}" if name else "") + f"\tProspectus · version {version_date(doc)}")
        r.font.size, r.font.color.rgb = Pt(8), muted
        hp.paragraph_format.tab_stops.add_tab_stop(tab, alignment=2)  # right
        fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        fp.text = ""

        def small(run: Any) -> None:
            run.font.size, run.font.color.rgb, run.bold = Pt(8), muted, False

        small(fp.add_run(f"{'Holding' if doc.get('scope') == 'holding' else 'Watch list'} · generated "
                         f"{str(doc.get('generated_at') or '')[:16]} UTC · advisory only\tPage "))
        _field(fp, "PAGE", small)
        small(fp.add_run(" / "))
        _field(fp, "NUMPAGES", small)
        fp.paragraph_format.tab_stops.add_tab_stop(tab, alignment=2)
    # images never wider than the text column
    max_w = d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin - Inches(0.3)
    for shape in d.inline_shapes:
        if shape.width and shape.width > max_w:
            ratio = max_w / shape.width
            shape.width, shape.height = int(max_w), int(shape.height * ratio)
    # repeat header rows; keep rows on one page
    for tbl in _all_tables(d):
        rows = tbl.rows
        if not rows:
            continue
        first = rows[0]._tr
        if first.xpath(".//w:tblHeader") or _row_is_header(first):
            trpr = first.get_or_add_trPr()
            if not trpr.xpath("./w:tblHeader"):
                el = OxmlElement("w:tblHeader")
                el.set(qn("w:val"), "true")
                trpr.append(el)
        for row in rows:
            if len(tbl.rows) > 1:
                trpr = row._tr.get_or_add_trPr()
                if not trpr.xpath("./w:cantSplit"):
                    trpr.append(OxmlElement("w:cantSplit"))
    font = str(s.get("docx_font") or "Calibri")
    _set_fonts(d.styles.element, font)
    _set_fonts(d.element, font)
    for sec in d.sections:
        _set_fonts(sec.header._element, font)
        _set_fonts(sec.footer._element, font)
    for st in d.styles:
        try:
            if st.type == 1 and st.font is not None:
                st.font.name = font
        except (AttributeError, ValueError):
            pass
    # a heading or chart caption never ends a page alone
    paras = d.paragraphs + [p for t in _all_tables(d) for row in t.rows for c in row.cells for p in c.paragraphs]
    for p in paras:
        nxt = p._p.getnext()
        if p.style is not None and str(p.style.name).lower().startswith("heading"):
            p.paragraph_format.keep_with_next = True
        elif nxt is not None and nxt.xpath(".//w:drawing") and not p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
    cp = d.core_properties
    cp.title = f"{sym} prospectus — version {version_date(doc)}"
    cp.subject = "Ticker Prospectus (advisory only)"
    cp.author = "Trade AI"
    cp.comments = ""
    buf = io.BytesIO()
    d.save(buf)
    return _font_table_sans(buf.getvalue(), font)


def _font_table_sans(data: bytes, font: str) -> bytes:
    """Mark the document font as a sans (swiss) face so a reader without it (LibreOffice on Linux, a phone viewer)
    substitutes a sans font as the HTML does, not a serif."""
    import zipfile  # noqa: PLC0415
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename == "word/fontTable.xml":
                text = blob.decode("utf-8")
                text = re.sub(r'(<w:font w:name="%s">)(.*?)(</w:font>)' % re.escape(font),
                              lambda m: m.group(1) + '<w:altName w:val="Carlito"/><w:charset w:val="00"/>'
                              '<w:family w:val="swiss"/><w:pitch w:val="variable"/>' + m.group(3), text, flags=re.S)
                blob = text.encode("utf-8")
            z.writestr(item, blob)
    return out.getvalue()


def _all_tables(d: Any) -> list:
    out, stack = [], list(d.tables)
    while stack:
        t = stack.pop(0)
        out.append(t)
        for row in t.rows:
            for cell in row.cells:
                stack.extend(cell.tables)
    return out


def _row_is_header(tr: Any) -> bool:
    """The first row of a data table converted from <thead>: every cell text upper-case and short."""
    texts = ["".join(t.text or "" for t in tc.iter() if t.tag.endswith("}t")) for tc in tr.iterchildren() if tc.tag.endswith("}tc")]
    return len(texts) >= 2 and all(x and x == x.upper() and len(x) <= 40 for x in texts)


# ── the hook ────────────────────────────────────────────────────────────────────────────────────────────────────────
def render_print(doc: dict, pngs: dict[str, bytes], *, fallback_pdf: Callable[[dict, dict], bytes | None],
                 fallback_docx: Callable[[dict, dict], bytes | None], cfg: dict | None = None,
                 report: dict | None = None) -> dict[str, bytes]:
    """{"prospectus.pdf": …, "prospectus.docx": …}: Chromium/LibreOffice when available, else the fallbacks.
    ``report`` (optional) receives {"pdf": "chromium"|"fallback", "docx": "libreoffice"|"fallback", "reasons": [...]}."""
    s = settings(cfg)
    rep = report if report is not None else {}
    rep.setdefault("reasons", [])
    files: dict[str, bytes] = {}
    pdf_bytes = docx_bytes = None
    sym = doc.get("symbol")
    if not s.get("enabled", True):
        rep["reasons"].append("print renderer disabled in config")
    else:
        try:
            with _host_lock(float(s["lock_wait_s"])) as got:
                if not got:
                    rep["reasons"].append(f"print lock busy > {s['lock_wait_s']}s")
                else:
                    pdf_bytes, docx_bytes = _render_locked(doc, s, rep)
        except Exception as exc:  # noqa: BLE001 - never fail the stored report over a download format
            rep["reasons"].append(f"print renderer error {type(exc).__name__}: {str(exc)[:200]}")
    rep["pdf"] = "chromium" if pdf_bytes else "fallback"
    rep["docx"] = "libreoffice" if docx_bytes else "fallback"
    if not pdf_bytes:
        pdf_bytes = fallback_pdf(doc, pngs)
    if not docx_bytes:
        docx_bytes = fallback_docx(doc, pngs)
    for reason in rep["reasons"]:
        LOG.warning("prospectus print %s: %s", sym, reason)
    if pdf_bytes:
        files["prospectus.pdf"] = pdf_bytes
    if docx_bytes:
        files["prospectus.docx"] = docx_bytes
    return files


def _render_locked(doc: dict, s: dict, rep: dict) -> tuple[bytes | None, bytes | None]:
    why = chromium_available()
    if why:
        rep["reasons"].append(f"pdf+docx fallback: {why}")
        return None, None
    pdf_bytes = docx_bytes = None
    with tempfile.TemporaryDirectory(prefix="prospectus-print-") as tmp:
        t = Path(tmp)
        page = t / "print.html"
        page.write_text(print_html(doc, paper=s["paper"]), encoding="utf-8")
        chart_ids = [c["id"] for c in doc.get("charts") or [] if c.get("id")]
        head, foot = _hf_templates(doc, s)
        m = s["margin_in"]
        job = {"_dir": tmp, "html_path": str(page), "pdf_path": str(t / "print.pdf"), "paper": s["paper"],
               "margin": {k: f"{v}in" for k, v in m.items()}, "header_html": head, "footer_html": foot,
               "shot_width_px": int(s["chart_width_px"]) + 64, "scale": s["chart_scale"],
               "shots": [{"selector": f"#chart-{cid} svg", "path": str(t / f"chart-{cid}.png")} for cid in chart_ids],
               "chromium_args": list(s.get("chromium_args") or []),
               "timeout_ms": int(float(s["chromium_timeout_s"]) * 1000)}
        res = run_chromium(job, float(s["chromium_timeout_s"]))
        if not res.get("ok"):
            rep["reasons"].append(f"pdf+docx fallback: {res.get('error')}")
            return None, None
        pdf_path = t / "print.pdf"
        if pdf_path.exists() and pdf_path.read_bytes()[:4] == b"%PDF":
            pdf_bytes = pdf_path.read_bytes()
        else:
            rep["reasons"].append("pdf fallback: chromium wrote no PDF")
        rep["blocked_requests"] = res.get("blocked_requests", 0)
        chart_files = {cid: (t / f"chart-{cid}.png").read_bytes() for cid in chart_ids if (t / f"chart-{cid}.png").exists()}
        try:
            if not soffice_path(s):
                raise RuntimeError("soffice not found")
            src = t / "prospectus.html"
            src.write_text(word_html(doc, chart_files, width_px=int(s["docx_image_width_px"])), encoding="utf-8")
            out = soffice_convert(src, t, "docx:MS Word 2007 XML", s, infilter="HTML (StarWriter)")
            docx_bytes = postprocess_docx(out.read_bytes(), doc, s)
            if not _docx_has_images(docx_bytes, len(chart_files)):
                raise RuntimeError("converted DOCX lost its chart images")
        except subprocess.TimeoutExpired:
            rep["reasons"].append(f"docx fallback: soffice timeout after {s['soffice_timeout_s']}s")
            docx_bytes = None
        except Exception as exc:  # noqa: BLE001
            rep["reasons"].append(f"docx fallback: {type(exc).__name__}: {str(exc)[:200]}")
            docx_bytes = None
    return pdf_bytes, docx_bytes


def _docx_has_images(data: bytes, want: int) -> bool:
    import zipfile  # noqa: PLC0415
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    return sum(1 for n in names if n.startswith("word/media/")) >= want


def data_uri(png: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(png).decode()
