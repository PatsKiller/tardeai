"""Prospectus institutional ("Street") blocks and the validation badge (operator 2026-10-27: "does it also have analyst
reviews and what are analysts saying along with catalysts? … something I can use on the Goldman desk, Citadel,
anywhere" and "We don't need to see [citations] on the frontend … If it's been validated, that's good.").

Synthetic tickers, names and numbers only (public repo); no database, no network, nothing written outside tmp.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
import io
import json
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from scripts.lib.prospectus import presentation as pr  # noqa: E402
from scripts.lib.prospectus import quality_gate as qg  # noqa: E402
from scripts.lib.prospectus import render, street, tables  # noqa: E402
from scripts.lib.prospectus.blocks import is_session  # noqa: E402
from scripts.lib.prospectus.inputs import Bundle  # noqa: E402

NOW = datetime(2026, 10, 26, 21, 0, tzinfo=timezone.utc)
SYM, NAME = "ZQXA", "Zeeco Synthetic Corp"


def sessions(start: date, end: date) -> list[str]:
    out, d = [], start
    while d <= end:
        if is_session(d):
            out.append(d.isoformat())
        d += timedelta(days=1)
    return out


def synthetic_bars() -> list[dict]:
    days = sessions(date(2022, 1, 3), date(2026, 10, 26))
    out = []
    for i, d in enumerate(days):
        px = 40.0 * (1.0005 ** i)
        if d == "2026-07-30":  # the session after the 2026-07-29 after-close report
            px *= 1.06
        out.append({"price_date": d, "close": round(px, 4)})
    return out


def quarter_reports() -> list[dict]:
    """Ten reported quarters (after the close) and one upcoming report."""
    rows = [{"Earnings Date": "2026-11-05T16:05:00-05:00", "EPS Estimate": 1.30, "Reported EPS": None, "Surprise(%)": None}]
    d = date(2026, 7, 29)
    for k in range(10):
        est = round(1.20 - 0.05 * k, 2)
        act = round(est * (1.04 if k != 3 else 0.97), 2)
        rows.append({"Earnings Date": f"{d.isoformat()}T16:05:00-04:00", "EPS Estimate": est, "Reported EPS": act,
                     "Surprise(%)": round((act - est) / est * 100, 2)})
        d -= timedelta(days=91)
    return rows


def frame_rows(rows: list[dict], index_key: str) -> list[dict]:
    """Same shape pro_analyst_fetch.frame_records stores: the frame index under ``index``."""
    return [{"index": r[index_key], **{k: v for k, v in r.items() if k != index_key}} for r in rows]


def synthetic_raw() -> dict:
    targets = []
    d = date(2026, 6, 28)
    i = 0
    while d <= date(2026, 10, 26):
        targets.append({"date": d.isoformat(), "mean": round(50 + i * 0.08, 2), "median": round(50 + i * 0.08, 2),
                        "low": 40.0, "high": 70.0 + (5 if i > 60 else 0), "rating_mean": 1.9, "rating_key": "buy",
                        "analysts": 21, "price": None})
        d += timedelta(days=1)
        i += 1
    ext = {
        "schema": "YahooAnalystExtended@v1", "snapshot_date": "2026-10-26",
        "upgrades_downgrades": [
            {"index": "2026-10-20T00:00:00", "Firm": "Example Securities", "ToGrade": "Buy", "FromGrade": "Hold",
             "Action": "up", "priceTargetAction": "Raises", "currentPriceTarget": 66.0, "priorPriceTarget": 55.0},
            {"index": "2026-09-02T00:00:00", "Firm": "Sample Capital", "ToGrade": "Neutral", "FromGrade": "Neutral",
             "Action": "main", "priceTargetAction": "Lowers", "currentPriceTarget": 48.0, "priorPriceTarget": 52.0},
            {"index": "2026-08-01T00:00:00", "Firm": "Demo Partners", "ToGrade": "Outperform", "FromGrade": "",
             "Action": "init", "priceTargetAction": "Announces", "currentPriceTarget": 62.0, "priorPriceTarget": 0.0},
        ],
        "earnings_dates": frame_rows(quarter_reports(), "Earnings Date"),
        "eps_trend": [{"index": "0q", "current": 1.30, "7daysAgo": 1.29, "30daysAgo": 1.27, "60daysAgo": 1.26, "90daysAgo": 1.25},
                      {"index": "0y", "current": 5.00, "7daysAgo": 4.98, "30daysAgo": 4.90, "60daysAgo": 4.85, "90daysAgo": 4.80},
                      {"index": "+1y", "current": 5.60, "7daysAgo": 5.60, "30daysAgo": 5.50, "60daysAgo": 5.45, "90daysAgo": 5.40}],
        "eps_revisions": [{"index": "0q", "upLast7days": 1, "upLast30days": 4, "downLast30days": 1, "downLast7Days": 0},
                          {"index": "0y", "upLast7days": 2, "upLast30days": 6, "downLast30days": 2, "downLast7Days": 0}],
        "earnings_estimate": [{"index": "0q", "avg": 1.30, "low": 1.20, "high": 1.40, "yearAgoEps": 1.10, "numberOfAnalysts": 18, "growth": 0.18},
                              {"index": "0y", "avg": 5.00, "low": 4.70, "high": 5.30, "yearAgoEps": 4.30, "numberOfAnalysts": 20, "growth": 0.16},
                              {"index": "+1y", "avg": 5.60, "low": 5.00, "high": 6.20, "yearAgoEps": 5.00, "numberOfAnalysts": 19, "growth": 0.12}],
        "revenue_estimate": [{"index": "0y", "avg": 21.0e9, "low": 20.0e9, "high": 22.0e9, "numberOfAnalysts": 17, "yearAgoRevenue": 18.5e9, "growth": 0.135}],
        "recommendations": [{"index": 0, "period": "0m", "strongBuy": 6, "buy": 10, "hold": 4, "sell": 1, "strongSell": 0},
                            {"index": 1, "period": "-1m", "strongBuy": 6, "buy": 9, "hold": 5, "sell": 1, "strongSell": 0},
                            {"index": 2, "period": "-2m", "strongBuy": 5, "buy": 9, "hold": 6, "sell": 1, "strongSell": 0},
                            {"index": 3, "period": "-3m", "strongBuy": 5, "buy": 8, "hold": 7, "sell": 1, "strongSell": 0}],
        "errors": {},
    }
    snaps = []
    for k, d in enumerate(sessions(date(2026, 6, 1), date(2026, 10, 26))):
        snaps.append({"date": d, "atm_iv": 30 + (k % 20), "pc_oi_ratio": 0.7 + k * 0.001, "pc_vol_ratio": 0.9,
                      "skew25": 3.5, "put_oi": 1000, "call_oi": 1400})
    xbrl = []
    for q, (st, en, rev, opi, ni) in enumerate([("2026-04-01", "2026-06-30", 5.4e9, 1.4e9, 1.0e9),
                                                 ("2026-01-01", "2026-03-31", 5.1e9, 1.3e9, 0.95e9),
                                                 ("2025-07-01", "2025-09-30", 4.9e9, 1.2e9, 0.9e9),
                                                 ("2025-04-01", "2025-06-30", 4.6e9, 1.1e9, 0.8e9)]):
        for m, v in (("revenue", rev), ("operating_income", opi), ("net_income", ni)):
            xbrl.append({"metric": m, "value": v, "start": st, "end": en, "filed": en[:8] + "28", "form": "10-Q",
                         "url": "https://www.sec.gov/Archives/edgar/data/0/synthetic-index.htm"})
    # Q4 2025 only exists inside the fiscal year: FY minus the three stated quarters.
    xbrl += [{"metric": "revenue", "value": 19.3e9, "start": "2025-01-01", "end": "2025-12-31", "filed": "2026-02-15", "form": "10-K", "url": None},
             {"metric": "revenue", "value": 4.5e9, "start": "2025-01-01", "end": "2025-03-31", "filed": "2025-04-28", "form": "10-Q", "url": None},
             {"metric": "cash", "value": 3.0e9, "start": None, "end": "2026-06-30", "filed": "2026-07-28", "form": "10-Q", "url": None},
             {"metric": "long_term_debt", "value": 5.5e9, "start": None, "end": "2026-06-30", "filed": "2026-07-28", "form": "10-Q", "url": None}]
    return {
        "contract": "StreetDetail@v1", "symbol": SYM, "errors": {}, "as_of": NOW.isoformat(),
        "target_history": targets, "extended": ext, "rating_history": [],
        "analyst_catalysts": [
            {"type": "analyst_upgrade", "headline": "Other Bank upgrades Zeeco Synthetic to Buy, price target $68",
             "source": "google_news:Example", "url": "https://news.example.invalid/a", "at": "2026-10-10T12:00:00+00:00"},
            {"type": "analyst_upgrade", "headline": "Example Securities upgrades Zeeco Synthetic to Buy",  # dup of yahoo row
             "source": "google_news:Example", "url": "https://news.example.invalid/b", "at": "2026-10-21T12:00:00+00:00"},
            {"type": "analyst_upgrade", "headline": "Financial stocks analyst roundup: what to watch",  # names no firm action
             "source": "google_news:Example", "url": "https://news.example.invalid/c", "at": "2026-10-11T12:00:00+00:00"},
        ],
        "material_catalysts": [{"type": "product_launch", "headline": "Zeeco Synthetic launches a synthetic widget",
                                "severity": "high", "source": "feed", "url": None, "at": "2026-09-15T13:00:00+00:00"}],
        "option_snapshots": snaps,
        "term_structure": {"captured_at": "2026-10-26T12:00:00+00:00", "front_iv_pct": 48.0, "back_iv_pct": 35.0,
                           "term_slope_pct": 13.0,
                           "term_structure": [{"exp": "2026-10-30", "atm_iv_pct": 40.0}, {"exp": "2026-11-06", "atm_iv_pct": 48.0},
                                              {"exp": "2026-11-20", "atm_iv_pct": 38.0}]},
        "iv_history": [], "xbrl": xbrl,
        "form4": {"rows": [{"filer_name": "Synthetic Officer", "filer_relation": "Officer", "transaction_type": "S-Sale",
                            "shares": 1000, "price": 60, "total_value": 60000, "filing_date": "2026-10-01", "sec_url": None},
                           {"filer_name": "Synthetic Director", "filer_relation": "Director", "transaction_type": "P-Purchase",
                            "shares": 500, "price": 58, "total_value": 29000, "filing_date": "2026-09-20", "sec_url": None}],
                  "symbol_rows_ever": 2, "store_symbols": 10, "store_last": "2026-10-23"},
        "holders_13f": [],
        "dividends": {"history": [{"ex_date": "2026-08-14", "amount": 0.25}, {"ex_date": "2026-05-15", "amount": 0.25},
                                  {"ex_date": "2026-02-13", "amount": 0.24}, {"ex_date": "2025-11-14", "amount": 0.24}],
                      "recorded_next_ex": None},
    }


class FakeSources:
    def __init__(self, raw=None, bars=None):
        self.raw = raw if raw is not None else synthetic_raw()
        self._bars = bars if bars is not None else synthetic_bars()

    def street(self, sym):
        return self.raw

    def bars(self, sym, days=1100):
        return self._bars


def bundle(fund: bool = False) -> Bundle:
    b = Bundle(symbol=SYM, scope="holding", name=NAME)
    b.add("valuation", "Forward P/E", value=11.5, source="Finviz enrichment cache", as_of="2026-10-26")
    b.add("analyst", "Analyst target mean", value=59.6, unit="$", source="yahoo_analyst_targets_history", as_of="2026-10-26")
    b.blocks["estimates"] = {"eps_surprises": [], "revisions_90d": {"up": None, "down": None}, "next_report": "2026-11-05",
                             "gaps": ["EPS surprise history: the house keeps the latest quarter only (0 of 4)",
                                      "90-day estimate revisions: not in the house"]}
    b.blocks["ownership_flows"] = {"short_float_pct": 2.1, "days_to_cover": 1.8, "institutional_pct": 70.0,
                                   "institutional_change_pct": 0.4, "insider_buys_90d": None, "insider_sells_90d": None}
    b.gaps += ["EPS surprise history: the house keeps the latest quarter only (0 of 4)", "90-day estimate revisions: not in the house"]
    b.context["next_earnings_date"] = "2026-11-05"
    return b


@pytest.fixture()
def built() -> Bundle:
    b = bundle()
    street.attach(FakeSources(), b, now=NOW, root=ROOT)
    return b


# ── analyst coverage ──────────────────────────────────────────────────────────────────────────────────────────
def test_rating_distribution_and_trend(built):
    dist = built.blocks["analyst_coverage"]["distribution"]
    assert dist["current"]["total"] == 21 and dist["current"]["bullish_pct"] == 76.2
    assert dist["three_months_ago"]["bullish_pct"] == 61.9 and dist["trend_90d"] == "more bullish"
    assert dist["bullish_change_pts"] == 14.3 and "stale" not in dist


def test_target_trend_uses_rows_near_each_horizon(built):
    tg = built.blocks["analyst_coverage"]["targets"]
    assert tg["mean"] == 59.6 and tg["analysts"] == 21
    assert tg["trend"]["d30"]["date"] == "2026-09-26" and tg["trend"]["d90"]["date"] == "2026-07-28"
    assert tg["trend"]["d30"]["mean_change_pct"] == pytest.approx((59.6 / tg["trend"]["d30"]["mean"] - 1) * 100, abs=0.01)


def test_target_trend_never_passes_an_old_row_off_as_30_days_ago():
    raw = synthetic_raw()
    raw["target_history"] = [r for r in raw["target_history"] if not "2026-08-15" <= r["date"] <= "2026-10-20"]
    b = bundle()
    street.attach(FakeSources(raw=raw), b, now=NOW, root=ROOT)
    assert b.blocks["analyst_coverage"]["targets"]["trend"]["d30"] is None


def test_action_log_structured_first_news_deduped_and_issuer_named(built):
    acts = built.blocks["analyst_coverage"]["actions"]
    firms = [(a["date"], a["firm"], a["action"]) for a in acts]
    assert firms == [("2026-10-20", "Example Securities", "up"), ("2026-10-10", "Other Bank", "up"),
                     ("2026-09-02", "Sample Capital", "main"), ("2026-08-01", "Demo Partners", "init")]
    up = acts[0]
    assert (up["from_grade"], up["to_grade"], up["prior_target"], up["target"]) == ("Hold", "Buy", 55.0, 66.0)
    news = acts[1]
    assert news["to_grade"] == "Buy" and news["target"] == 68.0 and news["url"].startswith("https://")
    assert built.blocks["analyst_coverage"]["action_sources"] == {"structured": 3, "from_news": 1, "news_rejected": 1}
    t90 = built.blocks["analyst_coverage"]["actions_90d"]
    assert t90["upgrades"] == 2 and t90["target_raises"] == 1 and t90["target_cuts"] == 1


def test_parse_action_headline():
    p = street.parse_action_headline("Big Firm & Co lowers Zeeco Synthetic price target to $45")
    assert p["firm"] == "Big Firm & Co" and p["action"] == "target_down" and p["target"] == 45.0
    assert street.parse_action_headline("Analyst roundup for the week")["firm"] is None


# ── estimates ─────────────────────────────────────────────────────────────────────────────────────────────────
def test_estimates_consensus_revisions_and_eight_quarters(built):
    est = built.blocks["estimates"]
    fy1 = next(c for c in est["consensus"] if c["period"] == "0y")
    assert (fy1["eps_avg"], fy1["eps_n"], fy1["rev_avg"], fy1["rev_n"], fy1["eps_growth_pct"]) == (5.0, 20, 21.0e9, 17, 16.0)
    tr = next(t for t in est["eps_trend"] if t["period"] == "0y")
    assert tr["chg_d30_pct"] == 2.04 and tr["chg_d90_pct"] == 4.17
    assert est["revisions_90d"] == {"up": 10, "down": 3, "window": "30d (current quarter + FY1)"}
    assert len(est["eps_surprises"]) == 8 and est["beat_rate"]["beats"] == 7 and est["beat_rate"]["misses"] == 1
    q = est["eps_surprises"][0]
    assert q["date"] == "2026-07-29" and q["timing"] == "amc" and q["base_date"] == "2026-07-29"
    assert q["one_day_pct"] == pytest.approx(6.05, abs=0.02)  # the after-close report moves the NEXT session
    assert not any(g.startswith(("EPS surprise history", "90-day estimate revisions")) for g in built.gaps)


def test_reaction_timing_rules():
    closes = street.Closes([{"price_date": d, "close": c} for d, c in
                            (("2026-03-02", 10), ("2026-03-03", 11), ("2026-03-04", 12), ("2026-03-05", 13),
                             ("2026-03-06", 14), ("2026-03-09", 15), ("2026-03-10", 16))])
    amc = street.reaction(closes, "2026-03-03", "amc")
    assert amc["base_date"] == "2026-03-03" and amc["one_day_pct"] == pytest.approx(9.09, abs=0.01)
    bmo = street.reaction(closes, "2026-03-03", "bmo")
    assert bmo["base_date"] == "2026-03-02" and bmo["one_day_pct"] == 10.0 and bmo["five_day_pct"] == 50.0


# ── catalysts ─────────────────────────────────────────────────────────────────────────────────────────────────
def test_catalyst_calendar_forward_and_backward(built):
    cal = built.blocks["catalyst_calendar"]
    fwd = {e["kind"]: e for e in cal["forward"]}
    er = fwd["earnings"]
    assert er["date"] == "2026-11-05" and er["time"] == "after the close" and "1.30" in er["detail"]
    mv = er["implied_move"]
    assert mv["expiry"] == "2026-11-06" and mv["days"] == 11 and mv["move_pct"] == round(48.0 * (11 / 365) ** 0.5, 2)
    assert fwd["ex_dividend"]["date"] == "2026-11-13" and "estimated" in fwd["ex_dividend"]["label"].lower()
    assert any(e["kind"] == "macro" and "FOMC" in e["label"] for e in cal["forward"])
    assert [e["date"] for e in cal["forward"]] == sorted(e["date"] for e in cal["forward"])
    back = cal["backward"]
    assert back[0]["date"] == "2026-09-15" or back[0]["kind"] == "earnings"
    assert any(b["kind"] == "product_launch" and b["one_day_pct"] is not None for b in back)


# ── positioning ───────────────────────────────────────────────────────────────────────────────────────────────
def test_positioning_options_insider_and_gaps(built):
    pos = built.blocks["positioning"]
    opt = pos["options"]
    assert opt["iv_sessions"] >= 60 and opt["iv_rank"] is not None and 0 <= opt["iv_percentile"] <= 100
    assert opt["put_call_oi_30d_ago"] is not None and opt["term_slope_pct"] == 13.0
    ins = pos["insider"]
    assert (ins["buys_90d"], ins["sells_90d"], ins["net_value_90d"]) == (1, 1, -31000.0)
    assert built.blocks["ownership_flows"]["insider_buys_90d"] == 1
    assert any("13F" in g for g in pos["gaps"]) and any("retired" in g for g in pos["gaps"])


def test_positioning_iv_rank_needs_enough_sessions():
    raw = synthetic_raw()
    raw["option_snapshots"] = raw["option_snapshots"][-20:]
    b = bundle()
    street.attach(FakeSources(raw=raw), b, now=NOW, root=ROOT)
    assert b.blocks["positioning"]["options"]["iv_rank"] is None
    assert any(g.startswith("IV rank: only 20") for g in b.blocks["positioning"]["gaps"])


# ── valuation history and XBRL ────────────────────────────────────────────────────────────────────────────────
def test_valuation_history_trailing_pe_in_own_range(built):
    pe = built.blocks["valuation_history"]["pe"]
    assert pe["months"] >= 12 and pe["min"] <= pe["current"] <= pe["max"] and 0 <= pe["percentile"] <= 100
    assert "Yahoo earnings dates" in pe["basis"] and built.blocks["valuation_history"]["forward_pe"] == 11.5


def test_xbrl_quarters_margin_yoy_net_debt_and_derived_q4(built):
    fx = built.blocks["fundamentals_xbrl"]
    by = {q["period_end"]: q for q in fx["quarters"]}
    q2 = by["2026-06-30"]
    assert q2["operating_margin_pct"] == 25.9 and q2["net_debt"] == 2.5e9
    assert q2["revenue_yoy_pct"] == pytest.approx((5.4 / 4.6 - 1) * 100, abs=0.01)
    q4 = by["2025-12-31"]  # FY 19.3 − (4.5 + 4.6 + 4.9)
    assert q4["revenue"] == pytest.approx(5.3e9) and q4["derived"] == ["revenue"]


# ── funds, facts, gate ────────────────────────────────────────────────────────────────────────────────────────
def test_fund_hides_company_blocks_and_keeps_calendar_and_positioning():
    b = bundle()
    street.attach(FakeSources(), b, now=NOW, fund=True, root=ROOT)
    for k in street.FUND_HIDDEN:
        assert k not in b.blocks
    assert "catalyst_calendar" in b.blocks and b.blocks["positioning"]["insider"] is None
    assert not any(g.startswith(("estimate revisions", "EPS / revenue consensus")) for g in b.gaps)


def test_writer_facts_are_cited_and_compact(built):
    labels = [f["label"] for f in built.facts]
    for want in ("Analysts rating buy or strong buy %", "Mean price target change over 30 days %",
                 "EPS consensus, Current fiscal year (FY1)", "EPS beats in the last 8 quarters",
                 "Upcoming: Earnings report", "Put/call ratio (open interest)", "Trailing P/E (reported EPS)",
                 "Insider net buying, 90 days (Form 4, $)"):
        assert want in labels, want
    assert any(lab.startswith("Analyst action: Example Securities — Upgrade") for lab in labels)
    assert all(f["id"].startswith("F") and f["source"] for f in built.facts)
    assert len(built.facts) < 90


def test_gate_street_checks_pass_on_good_blocks_and_flag_bad_ones(built):
    packet = {"symbol": SYM, "facts": built.facts, "blocks": built.blocks, "as_of": NOW.isoformat()}
    assert qg.check_street(packet, qg.load_settings()) == []
    bad = json.loads(json.dumps(built.blocks, default=str))
    bad["positioning"]["options"]["atm_iv_pct"] = 0.31  # a fraction recorded as %
    bad["catalyst_calendar"]["forward"][0]["implied_move"] = {"move_pct": 140.0, "expiry": "2026-11-06"}
    bad["analyst_coverage"]["targets"]["as_of"] = "2026-09-01"
    bad["estimates"]["eps_surprises"][0]["surprise_pct"] = -3.0
    codes = {f["code"] for f in qg.check_street({**packet, "blocks": bad}, qg.load_settings())}
    assert {"street_options_units", "street_implied_move_units", "street_targets_stale", "street_surprise_sign"} <= codes
    assert all(f["severity"] == qg.SOFT for f in qg.check_street({**packet, "blocks": bad}, qg.load_settings()))


# ── rendering: tables, no visible citations, one validation statement ─────────────────────────────────────────
def report_doc(built, *, gate_status="pass", unmatched=0) -> dict:
    return {"schema": "Prospectus@v1", "symbol": SYM, "name": NAME, "scope": "holding", "version": "20261026T210000Z",
            "generated_at": NOW.isoformat(), "model": "claude-opus-5-5", "status": "ready",
            "verdict_line": "HOLD — the Street is constructive [F2].",
            "sections": [{"id": "summary_verdict", "title": "Summary verdict",
                          "body_md": "Consensus mean target is $59.60 [F2]; forward P/E 11.5 [F1, W1].", "citations": []},
                         {"id": "analyst_view", "title": "Analyst view",
                          "body_md": "Example Securities moved to Buy [F2].", "citations": []}],
            "charts": [], "blocks": built.blocks, "facts": built.facts, "gaps": [],
            "gate": {"status": gate_status, "checked_at": "2026-10-26T20:59:00+00:00",
                     "reasons": [] if gate_status == "pass" else ["market_cap_units: market cap labelled $B exceeds the plausible maximum"],
                     "unreliable": {} if gate_status == "pass" else {"F1": "market_cap_units"},
                     "layers": {"post_write": {"unmatched_count": unmatched}}},
            "advisory": "READ_ONLY_ADVISORY"}


def test_street_tables_render_in_order(built):
    ids = [t["id"] for t in tables.street_tables(report_doc(built))]
    assert ids == list(tables.STREET_TABLE_IDS)
    acts = next(t for t in tables.street_tables(report_doc(built)) if t["id"] == "analyst_actions")
    assert acts["rows"][0][:5] == ["2026-10-20", "Example Securities", "Upgrade", "Hold → Buy", "$55.00 → $66.00"]
    fund = tables.street_tables(report_doc(built), hidden=pr.FUND_HIDDEN_BLOCKS)
    assert {t["id"] for t in fund} == {"catalyst_calendar", "catalyst_history", "positioning"}


def test_validation_line_validated_and_partial(built):
    v = pr.validation(report_doc(built))
    assert v["status"] == "validated" and v["line"].startswith("✓ Validated — 2 figures checked against 2 sources")
    assert v["by_source"] == {"Finviz enrichment cache": 1, "yahoo_analyst_targets_history": 1}
    p = pr.validation(report_doc(built, gate_status="soft_fail", unmatched=2))
    assert p["status"] == "partial" and p["not_validated"] == 3
    assert p["line"].startswith("⚠ Partial — 3 figures not validated: market cap labelled $B")
    assert p["unreliable"] == [{"id": "F1", "label": "Forward P/E", "reason": "market cap units"}]
    assert pr.validation({**report_doc(built), "gate": None})["status"] == "unchecked"


def test_html_has_no_visible_citations_and_one_validation_statement(built):
    doc = report_doc(built)
    html = render.html(doc)
    body = re.sub(r"<[^>]+>", " ", html.split("<body>")[1])
    assert "[F" not in body and "[W" not in body and not re.search(r"\bF\d+\b", body)
    assert "<sup" not in html and 'id="sources"' not in html and "Sources (" not in html
    assert html.count('data-validation="validated"') == 1 and "✓ Validated — 2 figures" in html
    assert 'data-fact-id="F1 W1" hidden' in html  # ids kept as data
    for title in ("Analyst coverage · what the Street says", "Analyst action log", "Estimates · consensus and revisions",
                  "Earnings history · beat / miss", "Catalyst calendar · next 120 days", "Positioning",
                  "Valuation vs its own history", "Reported quarters · SEC XBRL"):
        assert title in html, title
    assert "forward P/E 11.5." in re.sub(r"<[^>]+>", "", html)  # no space or bracket left behind


def test_pdf_print_html_has_validation_line_and_no_footnotes(built):
    html = render.html(report_doc(built, gate_status="soft_fail"), png_charts={}, theme="light")
    assert '<p class="pc-valid pc-valid-warn" data-validation="partial">⚠ Partial' in html
    assert "<sup" not in html and "pc-appendix" not in html and "[F" not in re.sub(r"<[^>]+>", "", html)


def test_docx_has_validation_line_no_superscripts_no_sources(built):
    pytest.importorskip("docx")
    from docx import Document
    data = render.docx(report_doc(built), {})
    d = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "✓ Validated — 2 figures checked against 2 sources" in text
    assert "[F" not in text + cells and "[W" not in text + cells
    assert not any(h.text == "Sources" for h in d.paragraphs if h.style.name.startswith("Heading"))
    assert not any(r.font.superscript for p in d.paragraphs for r in p.runs)
    assert "Example Securities" in cells and "Analyst action log" in text


def test_presentation_carries_validation_and_street(built):
    pres = pr.present(report_doc(built))
    assert pres["validation"]["status"] == "validated" and [t["id"] for t in pres["street"]][0] == "analyst_coverage"


# ── the writer of the extended capture (pro_analyst_fetch) ────────────────────────────────────────────────────
def test_extended_snapshot_shapes_and_breaker():
    pd = pytest.importorskip("pandas")
    import importlib.util
    import types
    from unittest import mock
    orig = Path.read_text

    def read_text(self, *a, **k):  # the module reads .env at import; no real file is read here
        return "" if self.name == ".env" else orig(self, *a, **k)

    spec = importlib.util.spec_from_file_location("paf_street", ROOT / "scripts" / "pro_analyst_fetch.py")
    mod = importlib.util.module_from_spec(spec)
    stub = {} if "psycopg2" in sys.modules else {"psycopg2": types.ModuleType("psycopg2")}  # CI has no psycopg2
    with mock.patch.object(Path, "read_text", read_text), mock.patch.dict(sys.modules, stub):
        spec.loader.exec_module(mod)

    class Tk:
        upgrades_downgrades = pd.DataFrame({"Firm": ["A", "B"], "ToGrade": ["Buy", "Hold"], "FromGrade": ["Hold", "Buy"],
                                            "Action": ["up", "down"]},
                                           index=pd.Index(pd.to_datetime(["2026-01-01", "2026-03-01"]), name="GradeDate"))
        eps_trend = pd.DataFrame({"current": [1.0, float("nan")]}, index=pd.Index(["0q", "0y"], name="period"))

        @property
        def earnings_history(self):
            raise RuntimeError("Too Many Requests. Rate limited.")

        def get_earnings_dates(self, limit=12):
            return pd.DataFrame({"EPS Estimate": [1.0]}, index=pd.Index(pd.to_datetime(["2026-04-01 16:00"]), name="Earnings Date"))

    out = mod.extended_snapshot(Tk(), attrs=(("upgrades_downgrades", 5), ("eps_trend", 5), ("earnings_history", 5)))
    assert [r["Firm"] for r in out["upgrades_downgrades"]] == ["B", "A"]  # newest first
    assert out["upgrades_downgrades"][0]["index"].startswith("2026-03-01")
    assert out["eps_trend"][1]["current"] is None and out["eps_trend"][0]["index"] == "0q"
    assert out["rate_limited"] is True and "earnings_history" in out["errors"]
    assert out["earnings_dates"][0]["index"].startswith("2026-04-01")
    json.dumps(out)  # JSON-safe for the data column


def test_e2e_street_fixture_in_sync(built):
    data = json.loads((ROOT / "apps/command-center-v3/e2e/fixtures/prospectus-street.json").read_text(encoding="utf-8"))
    pres = pr.present(report_doc(built))
    assert data["street"] == json.loads(json.dumps(pres["street"])), "regenerate e2e/fixtures/prospectus-street.json"
    assert data["validation"] == json.loads(json.dumps(pres["validation"]))


# ── coordinator review 2026-10-27 (print proof): thin EPS chart, peer credibility, pipe tables ────────────────────
def test_eps_surprise_chart_needs_three_quarters():
    one = {"id": "eps_surprise", "series": [{"name": "Surprise %", "points": [["2026-07-28", 2.7]]}]}
    three = {"id": "eps_surprise", "series": [{"name": "Surprise %", "points": [["a", 1.0], ["b", 2.0], ["c", -1.0]]}]}
    assert render.drawable(one) is False and render.drawable(three) is True
    html = render.html({"symbol": SYM, "charts": [one], "sections": [], "blocks": {}, "facts": []})
    assert 'id="chart-eps_surprise"' not in html


def test_peers_are_size_matched_and_mostly_filled():
    from scripts.lib.prospectus import blocks as bk

    def rec(ind, sector, cap_m, filled=True):
        base = {"industry": ind, "sector": sector, "market_cap_b": cap_m}
        if filled:
            base.update({"forward_pe": 20, "sales_growth_past_5y_pct": 5, "oper_margin_pct": 20, "roe_pct": 15,
                         "div_yield_pct": 1, "total_debt_equity": 0.5, "perf_year_pct": 10})
        return base
    cache = {"ZBIG": rec("Gadgets", "Tech", 3_000_000), "ZTNY": rec("Gadgets", "Tech", 2_000),  # 1/1500th: no peer
             "ZMID": rec("Gadgets", "Tech", 900_000), "ZEMP": rec("Gadgets", "Tech", 2_500_000, filled=False),
             "ZSW1": rec("Software", "Tech", 2_800_000), "ZSW2": rec("Software", "Tech", 1_500_000),
             "ZOTH": rec("Banks", "Financial", 2_900_000)}
    p = bk.peer_comparison("ZBIG", cache)
    syms = [r["symbol"] for r in p["rows"]]
    assert syms[0] == "ZBIG" and "ZTNY" not in syms and "ZEMP" not in syms and "ZOTH" not in syms
    assert set(syms[1:]) == {"ZMID", "ZSW1", "ZSW2"}
    assert any(g.startswith("peers widened to the Tech sector by size") for g in p["gaps"])


def test_markdown_pipe_tables_render_as_tables():
    md = "Intro.\n| Metric | Value |\n|---|---:|\n| Forward P/E | 11.5 [F1] |\n| Yield | 0.7% |\nAfter."
    html = render.md_to_html(md, {})
    assert "<table" in html and "<th>Metric</th>" in html and "| Metric" not in html and "[F1]" not in html
    assert pr.md_segments(md)[1][0] == "table"


def test_docx_pipe_table_is_a_word_table(built):
    pytest.importorskip("docx")
    from docx import Document
    doc = report_doc(built)
    doc["sections"][1]["body_md"] = "| Firm | Action |\n|---|---|\n| Example Securities | Upgrade [F2] |"
    d = Document(io.BytesIO(render.docx(doc, {})))
    text = "\n".join(p.text for p in d.paragraphs)
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "| Firm" not in text and "Upgrade" in cells and "[F2]" not in " ".join(cells)
